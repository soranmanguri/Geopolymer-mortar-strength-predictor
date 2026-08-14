"""
Summary tables for the paper.

Table 1 — per-reference range summary (one row per source study).
Table 2 — descriptive statistics per input/output variable.
"""

from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _load_raw(file_path: str) -> pd.DataFrame:
    df = pd.read_excel(file_path, sheet_name="Sheet1")
    df.columns = df.columns.str.strip()
    return df


def _format_range(series: pd.Series, decimals: int = 2) -> str:
    s = pd.to_numeric(series, errors="coerce").dropna()
    if len(s) == 0:
        return ""
    mn, mx = s.min(), s.max()
    if np.isclose(mn, mx):
        return str(int(mn)) if float(mn).is_integer() else f"{mn:.{decimals}f}".rstrip("0").rstrip(".")
    a = f"{mn:.{decimals}f}".rstrip("0").rstrip(".")
    b = f"{mx:.{decimals}f}".rstrip("0").rstrip(".")
    return f"{a}–{b}"


# ---------------------------------------------------------------------------
# Table 1
# ---------------------------------------------------------------------------

_T1_RENAME = {
    "Reference":           "Ref.",
    "Alk: Binder":         "l/b",
    "All Binders":         "B (kg/m3)",
    "Fine Aggregate":      "FA (kg/m3)",
    "NaOH":                "SH (kg/m3)",
    "Na2SiO3":             "SS (kg/m3)",
    "Molarity (M)":        "M",
    "Na2SiO3/NaOH":        "SS/SH",
    "Nano Silica (Kg)":    "nS (kg/m3)",
    "Curing Condition":    "T (°C)",
    "Comp.(MPa, 28 days)": "CS (MPa)",
}

_T1_SUMMARY_COLS = [
    "l/b", "B (kg/m3)", "FA (kg/m3)", "SH (kg/m3)", "SS (kg/m3)",
    "M", "SS/SH", "nS (kg/m3)", "T (°C)", "CS (MPa)",
]


def build_table1(
    file_path: str = "Optimization - Final.xlsx",
    output_path: str = "Table1_like_paper.xlsx",
) -> pd.DataFrame:
    """Build Table 1: per-reference value ranges + global Min/Max/St.Div."""
    df = _load_raw(file_path)
    keep = [c for c in _T1_RENAME if c in df.columns]
    df = df[keep].rename(columns=_T1_RENAME)

    rows = []
    for ref, grp in df.groupby("Ref."):
        row = {"Ref.": ref}
        for col in _T1_SUMMARY_COLS:
            row[col] = _format_range(grp[col]) if col in grp.columns else ""
        rows.append(row)

    table = pd.DataFrame(rows)

    for label, fn in [("Min.", "min"), ("Max.", "max"), ("St.Div.", "std")]:
        stat_row = {"Ref.": label}
        for col in _T1_SUMMARY_COLS:
            if col in df.columns:
                s = pd.to_numeric(df[col], errors="coerce").dropna()
                val = getattr(s, fn)()
                stat_row[col] = f"{val:.2f}".rstrip("0").rstrip(".")
            else:
                stat_row[col] = ""
        rows.append(stat_row)

    table = pd.DataFrame(rows)
    table.to_excel(output_path, index=False)
    print(f"Saved: {output_path}")
    print(table.to_string(index=False))
    return table


# ---------------------------------------------------------------------------
# Table 2
# ---------------------------------------------------------------------------

_T2_COL_MAP = {
    "l/b":         "Alk: Binder",
    "B (kg/m3)":   "All Binders",
    "Fly ash":     "Fly ash",
    "Slag":        "Slag",
    "Metakaolin":  "Metakaoline",
    "FA (kg/m3)":  "Fine Aggregate",
    "SH (kg/m3)":  "NaOH",
    "SS (kg/m3)":  "Na2SiO3",
    "M":           "Molarity (M)",
    "SS/SH":       "Na2SiO3/NaOH",
    "nS (kg/m3)":  "Nano Silica (Kg)",
    "T (°C)":      "Curing Condition",
    "CS (MPa)":    "Comp.(MPa, 28 days)",
}

_T2_ORDER = [
    "l/b", "B (kg/m3)", "Fly ash", "Slag", "Metakaolin",
    "FA (kg/m3)", "SH (kg/m3)", "SS (kg/m3)", "M", "SS/SH",
    "nS (kg/m3)", "T (°C)", "CS (MPa)",
]

_T2_BOLD = {"l/b", "B (kg/m3)", "FA (kg/m3)", "SH (kg/m3)", "SS (kg/m3)",
            "M", "SS/SH", "nS (kg/m3)", "T (°C)", "CS (MPa)"}
_T2_INDENT = {"Fly ash", "Slag", "Metakaolin"}


def build_table2(
    file_path: str = "Optimization - Final.xlsx",
    output_path: str = "Table2_with_binder_details.xlsx",
) -> pd.DataFrame:
    """Build Table 2: descriptive statistics per variable."""
    df = _load_raw(file_path)

    available = {k: v for k, v in _T2_COL_MAP.items() if v in df.columns}
    missing = [v for v in _T2_COL_MAP.values() if v not in df.columns]
    if missing:
        print(f"Note: columns not found in file (skipped): {missing}")

    data = df[list(available.values())].apply(pd.to_numeric, errors="coerce")
    data.columns = list(available.keys())

    table = pd.DataFrame({
        "No. of data": data.count(),
        "Average":     data.mean(),
        "Median":      data.median(),
        "St.Div.":     data.std(),
        "Min.":        data.min(),
        "Max.":        data.max(),
        "Variance":    data.var(),
        "Skewness":    data.skew(),
        "Kurtosis":    data.kurt(),
    }).round(2)

    table.insert(0, "Model parameters", table.index)
    table = table.reset_index(drop=True)

    order_map = {k: i for i, k in enumerate(_T2_ORDER)}
    table["_order"] = table["Model parameters"].map(order_map).fillna(99)
    table = table.sort_values("_order").drop(columns="_order").reset_index(drop=True)

    table.to_excel(output_path, index=False, engine="openpyxl")

    # Apply formatting
    wb = load_workbook(output_path)
    ws = wb.active
    for row in range(2, ws.max_row + 1):
        val = ws.cell(row=row, column=1).value
        if val in _T2_BOLD:
            ws.cell(row=row, column=1).font = Font(bold=True)
        elif val in _T2_INDENT:
            ws.cell(row=row, column=1).alignment = Alignment(indent=1)
    for col in ws.columns:
        letter = col[0].column_letter
        width = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[letter].width = width + 2
    wb.save(output_path)

    print(f"Saved: {output_path}")
    print(table.to_string(index=False))
    return table
